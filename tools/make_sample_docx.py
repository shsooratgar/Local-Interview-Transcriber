"""Generate a SYNTHETIC noScribe-style .docx to test the parser and report
(this is fake content, not real data). It mimics the real layout:

  - A meta table at the top (one row already has a Summary filled, to verify
    we DON'T overwrite it).
  - A transcript body with mixed speaker labels (names + S01-style), inline
    //interjections//, (.) pause markers, and a couple of planted errors.
"""
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "testsets" / "sample_transcript.docx"

TABLE_ROWS = [
    ("Meeting title", "R&D interview — catalysis & ML"),
    ("Date/ duration", "2026-06-10 / 45 min"),
    ("Participants", "All R&D invited"),
    ("Summary", "Existing notes (already filled — must NOT be overwritten)."),
    ("Reflections", ""),
]

# planted errors: "cattle is" -> catalyst ; "reinforced film learning" -> reinforcement learning
BODY = [
    "Alex: So could you walk me through the reactor setup? (.)",
    "robin: Sure, it's a fixed-bed reactor. (..) We rely on a cattle is to speed up the "
    "reaction without consuming it. //Alex: Okay.// And then we measure the conversion.",
    "S07: I can explain that part.",
    "Alex: Oh, okay.",
    "S05: We also trained a model. //SAM: Which one?// Or rather a few of them.",
    "SAM: Okay, then please continue.",
    "S05: Yeah, we used reinforced film learning to optimize the policy. "
    "//Alex: Interesting.// And then we validated it on held-out runs.",
    "Alex: Okay. Okay. (..) Thanks for explaining all of that.",
]


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, val in TABLE_ROWS:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = val
    doc.add_paragraph("")
    for line in BODY:
        doc.add_paragraph(line)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
