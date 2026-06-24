#!/usr/bin/env python3
"""Apply accepted review decisions to a COPY of the base .docx, producing a
corrected document. The original is never modified.

decisions map (from the viewer): { "<key>": {"s": "keep"|"drop"|"", "e": "<edit>"} }
where  c<i> -> review["disagreements"][i]
       r<i> -> review["recovered"][i]   (replace the dropped remnant)
       o<i> -> review["omissions"][i]   (insert the missing passage)

Replacements are targeted to the exact paragraph (via each flag's segment index),
falling back to a first-match search. Anything that can't be located is reported
as 'skipped' so you can do it by hand.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from transcript_io import _parse_turn


def _final_text(flag, dec):
    edit = (dec.get("e") or "").strip()
    return edit or (flag.get("llm") or {}).get("correction") or flag.get("audio_text") or ""


def _rewrite(p, new_full):
    """Replace a paragraph's text, keeping its style (collapses runs)."""
    if p.runs:
        p.runs[0].text = new_full
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(new_full)


def _dedup_boundary(before, new, after):
    """Drop a boundary word of `new` if it just repeats the adjacent transcript word."""
    nw, bw, aw = new.split(), before.split(), after.split()
    if nw and aw and nw[-1].lower().strip(".,!?") == aw[0].lower().strip(".,!?"):
        nw = nw[:-1]
    if nw and bw and nw[0].lower().strip(".,!?") == bw[-1].lower().strip(".,!?"):
        nw = nw[1:]
    return " ".join(nw)


def _replace_in_para(p, old, new):
    if not old or old not in p.text:
        return False
    full = p.text
    idx = full.find(old)
    before, after = full[:idx], full[idx + len(old):]
    _rewrite(p, before + _dedup_boundary(before, new, after) + after)
    return True


def _replace_first(paragraphs, old, new):
    for p in paragraphs:
        if _replace_in_para(p, old, new):
            return True
    return False


def _seg_to_para(doc):
    """Map each flattened base-segment index to its paragraph (same order as
    transcript_io parses the docx) so replacements hit the right spot."""
    mapping = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        for _ in _parse_turn(t):
            mapping.append(p)
    return mapping


def apply_decisions(base_docx, review, decisions, out_path):
    doc = Document(str(base_docx))
    paras = doc.paragraphs
    s2p = _seg_to_para(doc)
    counts = {"applied": 0, "skipped": 0}

    def do_replace(flag, dec):
        new = _final_text(flag, dec)
        seg = flag.get("seg")
        ok = False
        if new:
            if seg is not None and seg < len(s2p):
                ok = _replace_in_para(s2p[seg], flag.get("base_text", ""), new)
            if not ok:
                ok = _replace_first(paras, flag.get("base_text", ""), new)
        counts["applied" if ok else "skipped"] += 1

    for i, d in enumerate(review.get("disagreements", [])):
        dec = decisions.get(f"c{i}")
        if dec and dec.get("s") == "keep":
            do_replace(d, dec)

    for i, d in enumerate(review.get("recovered", [])):
        dec = decisions.get(f"r{i}")
        if dec and dec.get("s") == "keep":
            do_replace(d, dec)

    for i, o in enumerate(review.get("omissions", [])):
        dec = decisions.get(f"o{i}")
        if not dec or dec.get("s") != "keep":
            continue
        text = (dec.get("e") or "").strip() or o.get("audio_text", "")
        seg = o.get("seg")
        if text and seg is not None and seg < len(s2p):
            p = s2p[seg]
            _rewrite(p, (p.text + " " + text).strip())
            counts["applied"] += 1
        else:
            counts["skipped"] += 1

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {**counts, "out": str(out_path)}


if __name__ == "__main__":
    import sys
    review = json.loads(Path(sys.argv[2]).read_text())
    decisions = json.loads(Path(sys.argv[3]).read_text())
    print(apply_decisions(sys.argv[1], review, decisions, sys.argv[4]))
