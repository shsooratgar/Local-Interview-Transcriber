#!/usr/bin/env python3
"""Write an annotated .docx report from an Option-A review result.

We only ADD to the document; the original transcript text and table are left
intact. A clearly-labelled review section is appended at the end, with tables of
suggested corrections, recovered passages, possible omissions and not-heard text.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


def _heading(doc, text, level=1):
    """Add a heading WITHOUT depending on the built-in 'Heading N' styles.
    Some transcript .docx files (e.g. certain Word exports) ship with those
    styles stripped, and `doc.add_heading` then raises
    KeyError: "no style with name 'Heading 1'". Fall back to a bold, sized
    paragraph so the report still writes."""
    try:
        return doc.add_heading(text, level=level)
    except KeyError:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(15 if level <= 1 else 13)
        return p


def _mmss(s):
    if s is None:
        return "—"
    s = int(s)
    return f"{s // 60:02d}:{s % 60:02d}"


def _fill_cell_bullets(cell, items):
    cell.text = ""
    cell.paragraphs[0].text = ("• " + items[0]) if items else "(none)"
    for it in items[1:]:
        cell.add_paragraph("• " + it)


def _grid(doc, headers, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    try:
        t.style = "Table Grid"
    except Exception:
        pass
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
    return t


def write_docx_report(base_docx, review, out_path):
    doc = Document(str(base_docx))
    doc.add_page_break()
    _heading(doc, "LLM review — automatic suggestions, please verify", 1)
    doc.add_paragraph(f"(transcript/audio match coverage: {review.get('match_coverage', '?')}. "
                      f"Times refer to the audio; use the .html viewer to hear each one.)")

    _heading(doc, "Suggested corrections", 2)
    corr = review.get("disagreements", [])
    if corr:
        t = _grid(doc, ["Time", "Speaker", "In transcript", "Suggested (from audio)"])
        for d in corr:
            c = t.add_row().cells
            c[0].text = _mmss(d.get("start"))
            c[1].text = str(d.get("speaker") or "")
            c[2].text = d.get("base_text", "")
            c[3].text = (d.get("llm", {}).get("correction") or d.get("audio_text", ""))
    else:
        doc.add_paragraph("None found.")

    _heading(doc, "Recovered passages (transcript dropped these — often another speaker)", 2)
    rec = review.get("recovered", [])
    if rec:
        t = _grid(doc, ["Time", "Speaker", "In transcript", "Recovered (from audio)", "Other speaker?"])
        for d in rec:
            c = t.add_row().cells
            c[0].text = _mmss(d.get("start"))
            c[1].text = str(d.get("speaker") or "")
            c[2].text = d.get("base_text", "")
            c[3].text = (d.get("llm", {}).get("correction") or d.get("audio_text", ""))
            c[4].text = "yes" if d.get("other_speaker") else ""
    else:
        doc.add_paragraph("None found.")

    _heading(doc, "Possible omissions (heard in audio, missing from transcript)", 2)
    oms = review.get("omissions", [])
    if oms:
        t = _grid(doc, ["Time", "Heard in audio"])
        for o in oms:
            c = t.add_row().cells
            c[0].text = _mmss(o.get("start"))
            c[1].text = o.get("audio_text", "")
    else:
        doc.add_paragraph("None found.")

    _heading(doc, "Not heard in audio (in transcript only — possible fabrication, verify)", 2)
    bos = review.get("base_only", [])
    if bos:
        t = _grid(doc, ["~Time", "Speaker", "In transcript (never heard)"])
        for d in bos:
            c = t.add_row().cells
            c[0].text = _mmss(d.get("start"))
            c[1].text = str(d.get("speaker") or "")
            c[2].text = d.get("base_text", "")
    else:
        doc.add_paragraph("None found.")

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


if __name__ == "__main__":
    import json
    import sys
    review = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
    print("saved ->", write_docx_report(sys.argv[1], review, sys.argv[3]))
